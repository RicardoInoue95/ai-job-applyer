from pathlib import Path

from docx import Document

from jobapplier.resume_parser.exceptions import ExtractionError

from .base import BaseExtractor


class DOCXExtractor(BaseExtractor):
    def __init__(self, path: Path):
        super().__init__(path)

    def extract(self) -> str:
        try:
            doc = Document(self.path)
            paragraphs = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    paragraphs.append(text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        paragraphs.append(row_text)
            return "\n".join(paragraphs)
        except Exception as exc:
            raise ExtractionError(f"Erro ao extrair texto do DOCX: {exc}") from exc
